"""样式清理 module：接收已打开的文档，返回结构化清理结果。

不碰磁盘、不碰 GUI——保存与展示由调用方（GUI/CLI adapter 或 batch）负责。
"""
from dataclasses import dataclass
from enum import Enum, auto


class StyleCategory(Enum):
    """清理结果中的中性样式类别；中文展示名由 adapter 映射。"""
    PARAGRAPH = auto()
    CHARACTER = auto()
    TABLE = auto()
    OTHER = auto()


@dataclass
class DeletedStyle:
    name: str
    category: StyleCategory


def _category_of(style) -> StyleCategory:
    from docx.enum.style import WD_STYLE_TYPE
    mapping = {
        WD_STYLE_TYPE.PARAGRAPH: StyleCategory.PARAGRAPH,
        WD_STYLE_TYPE.CHARACTER: StyleCategory.CHARACTER,
        WD_STYLE_TYPE.TABLE: StyleCategory.TABLE,
    }
    try:
        return mapping.get(style.type, StyleCategory.OTHER)
    except Exception:
        return StyleCategory.OTHER


def _styles_in_paragraph(paragraph) -> set:
    """一个段落（含其字符 run）引用的样式名。"""
    names = set()
    if paragraph.style and paragraph.style.name:
        names.add(paragraph.style.name)
    for run in paragraph.runs:
        if run.style and run.style.name:
            names.add(run.style.name)
    return names


def _collect_used_style_names(document) -> set:
    """收集文档中引用的样式名。

    目前覆盖：正文段落、表格（含单元格段落与字符 run）、正文字符 run。
    尚未覆盖：页眉/页脚、脚注/尾注、文本框（见 CONTEXT.md 定义与 #4）。
    """
    used = set()

    for paragraph in document.paragraphs:
        used |= _styles_in_paragraph(paragraph)

    for table in document.tables:
        if table.style and table.style.name:
            used.add(table.style.name)
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    used |= _styles_in_paragraph(paragraph)

    return used


def clean_document(document) -> list:
    """删除文档中未使用的自定义样式，返回 DeletedStyle 列表。

    内置样式无论是否使用都永不删除。
    """
    used = _collect_used_style_names(document)

    to_delete = [
        style for style in document.styles
        if style.name not in used and not style.builtin
    ]

    deleted = []
    for style in to_delete:
        try:
            document.styles[style.name].delete()
            deleted.append(DeletedStyle(name=style.name, category=_category_of(style)))
        except Exception:
            # 某些样式可能无法删除，与既有行为一致：忽略
            continue
    return deleted
